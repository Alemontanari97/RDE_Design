# P1 TRACE step 1: full-text dump of the user's final deck (READ-ONLY).
from pptx import Presentation
import os

SRC = r"C:\Users\amont\Desktop\Presentazione_ESA.pptx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "trace", "ESA_FINAL_TEXT.md")
os.makedirs(os.path.dirname(OUT), exist_ok=True)


def walk(shape):
    if shape.shape_type == 6:
        for sub in shape.shapes:
            yield from walk(sub)
    elif getattr(shape, "has_text_frame", False):
        yield shape


p = Presentation(SRC)
lines = ["# ESA_FINAL_TEXT — dump testuale Presentazione_ESA.pptx "
         "(35 slide, salvata 2026-08-24; READ-ONLY probe per TRACE P1)", ""]
for i, s in enumerate(p.slides, 1):
    lines.append(f"\n## slide {i}")
    for top in s.shapes:
        for sh in walk(top):
            t = sh.text_frame.text.strip()
            if t and t not in ("Rotating Detonation Engine Activities",
                               "T(H)RUST team"):
                t = t.replace("\n", " | ")
                lines.append(f"- {t}")
    if s.has_notes_slide:
        nt = s.notes_slide.notes_text_frame.text.strip()
        if nt:
            lines.append(f"  [NOTE: {len(nt)} ch]")
with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines) + "\n")
print("wrote", OUT, len(lines), "lines")
