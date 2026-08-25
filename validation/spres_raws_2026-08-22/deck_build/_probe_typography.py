# S4 lot-2 probe (CKP-S3-5(i)): measure title + footer geometry/style on
# every slide of the built deck. Read-only; prints a per-slide table.
from pptx import Presentation
from pptx.util import Emu


def walk(shape):
    if shape.shape_type == 6:
        for sub in shape.shapes:
            yield from walk(sub)
    elif getattr(shape, "has_text_frame", False):
        yield shape


def fmt_in(v):
    return f"{Emu(v).inches:.2f}" if v is not None else "?"


def run_info(sh):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip():
                col = None
                try:
                    if r.font.color and r.font.color.type is not None:
                        col = str(r.font.color.rgb)
                except Exception:
                    pass
                sz = r.font.size.pt if r.font.size else None
                return sz, r.font.bold, col, p.alignment
    return None, None, None, None


p = Presentation("SPRES_deck_v1.pptx")
print("=== TITLES (top-of-slide text) ===")
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        is_title = False
        if sh.has_text_frame and sh.top is not None and Emu(sh.top).inches < 1.2:
            t = sh.text_frame.text.strip()
            if t and len(t) > 12:
                is_title = True
        if is_title:
            sz, bold, col, align = run_info(sh)
            kind = "PH" if sh.is_placeholder else "TB"
            print(f"s{i:02d} {kind} '{sh.name[:18]}' L={fmt_in(sh.left)} "
                  f"T={fmt_in(sh.top)} W={fmt_in(sh.width)} sz={sz} "
                  f"bold={bold} col={col} align={align} :: {t[:40]!r}")

print("\n=== FOOTER TEXTS ===")
for i, s in enumerate(p.slides, 1):
    for top_sh in s.shapes:
        for sh in walk(top_sh):
            t = sh.text_frame.text.strip()
            if t in ("Rotating Detonation Engine Activities", "T(H)RUST team"):
                sz, bold, col, align = run_info(sh)
                kind = "PH" if getattr(sh, "is_placeholder", False) else "TB"
                print(f"s{i:02d} {kind} '{sh.name[:22]}' L={fmt_in(sh.left)} "
                      f"T={fmt_in(sh.top)} W={fmt_in(sh.width)} sz={sz} "
                      f"bold={bold} col={col} align={align} :: {t[:20]!r}")
