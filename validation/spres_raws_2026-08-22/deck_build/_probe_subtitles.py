from pptx import Presentation
from pptx.util import Emu
p = Presentation("SPRES_deck_v1_new.pptx")
print("subtitle-zone runs (top<1.3, 24-31pt):")
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        if sh.top is None or Emu(sh.top).inches >= 1.3:
            continue
        for par in sh.text_frame.paragraphs:
            hit = None
            for r in par.runs:
                if r.text.strip() and r.font.size and 24 <= r.font.size.pt < 32:
                    col = None
                    try:
                        col = str(r.font.color.rgb)
                    except Exception:
                        pass
                    hit = (r.font.size.pt, r.font.bold, col, r.text[:38])
                    break
            if hit:
                print(f"s{i:02d}", hit)
            break
