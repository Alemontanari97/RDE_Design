from pptx import Presentation
from pptx.util import Emu


def walk(shape, depth=0):
    if shape.shape_type == 6:
        for sub in shape.shapes:
            yield from walk(sub, depth + 1)
    else:
        yield shape, depth


p = Presentation("SPRES_deck_v1_new.pptx")
for idx in (2, 41):  # s03 host, s42 new
    s = p.slides[idx]
    print(f"--- slide {idx+1}: bottom-zone shapes ---")
    for top_sh in s.shapes:
        for sh, d in walk(top_sh):
            if sh.top is None:
                continue
            t = Emu(sh.top).inches
            if t < 6.2:
                continue
            h = Emu(sh.height).inches if sh.height else 0
            w = Emu(sh.width).inches if sh.width else 0
            le = Emu(sh.left).inches if sh.left else 0
            fill = ""
            try:
                if sh.fill.type is not None and sh.fill.type == 1:
                    fill = str(sh.fill.fore_color.rgb)
            except Exception:
                pass
            txt = ""
            if getattr(sh, "has_text_frame", False):
                txt = sh.text_frame.text.strip()[:30]
                anch = sh.text_frame.vertical_anchor
            else:
                anch = None
            print(f"  d{d} '{sh.name[:24]}' L={le:.2f} T={t:.2f} W={w:.2f} "
                  f"H={h:.2f} fill={fill} anch={anch} :: {txt!r}")
