# S-PRES Block 0 — ppt_Heister template ROUND-TRIP fidelity test
# (carrier: ADVISORY_Spres_prompt addendum 2, "template round-trip
# fidelity test EARLY at Block 0"). python-pptx 1.0.2 pinned env.
#
# Question under test: does the Heister template survive
# open -> save (T0) and open -> add slide -> save (T1) without losing
# content parts (slides/media/masters/layouts/theme/notes)?
# REJECTOR: the test FAILS if any content part class is dropped or
# slide/master/layout/media counts change unexpectedly. It can reject:
# python-pptx is known to drop parts it does not model (e.g. some
# embedded objects); if that hits content, verdict = FAIL and the
# vehicle decision must be re-opened.
import io
import sys
import zipfile
from pptx import Presentation
from pptx.util import Emu

SRC = r"c:\Users\amont\Claude\Projects\Presentazione RDE CVA\rde-lecture-code\reference_presentations\ppt_Heister.pptx"
OUTDIR = r"C:\Users\amont\AppData\Local\Temp\claude\c--Users-amont-Claude-Projects-Presentazione-RDE-CVA-rde-lecture-code\dc66e145-ec4b-422c-bc42-0bf882c8a3a0\scratchpad\roundtrip"
REPORT = r"c:\Users\amont\Claude\Projects\Presentazione RDE CVA\rde-lecture-code\validation\spres_raws_2026-08-22\ROUNDTRIP_report.md"

import os
os.makedirs(OUTDIR, exist_ok=True)
T0 = os.path.join(OUTDIR, "roundtrip_0.pptx")
T1 = os.path.join(OUTDIR, "roundtrip_1.pptx")

def zparts(path):
    with zipfile.ZipFile(path) as z:
        return sorted(z.namelist())

def classify(name):
    if name.startswith("ppt/slides/") and name.endswith(".xml"):
        return "slide"
    if name.startswith("ppt/media/"):
        return "media"
    if name.startswith("ppt/slideMasters/") and name.endswith(".xml"):
        return "master"
    if name.startswith("ppt/slideLayouts/") and name.endswith(".xml"):
        return "layout"
    if name.startswith("ppt/theme/"):
        return "theme"
    if name.startswith("ppt/notesSlides/") and name.endswith(".xml"):
        return "notes"
    if name.startswith("ppt/embeddings/"):
        return "embedding"
    if name.startswith("ppt/charts/") or "/charts/" in name:
        return "chart"
    return "other"

def counts_by_class(names):
    c = {}
    for n in names:
        k = classify(n)
        c[k] = c.get(k, 0) + 1
    return c

def pptx_stats(path):
    prs = Presentation(path)
    return {
        "n_slides": len(prs.slides),
        "n_masters": len(prs.slide_masters),
        "n_layouts": sum(len(m.slide_layouts) for m in prs.slide_masters),
        "w_emu": int(prs.slide_width),
        "h_emu": int(prs.slide_height),
    }

lines = []
def log(s=""):
    lines.append(s)

log("# ROUND-TRIP fidelity report — ppt_Heister.pptx (S-PRES Block 0)")
log("")
log("Command of record: python validation/spres_raws_2026-08-22/roundtrip_test.py")
log("Env: python-pptx 1.0.2 (pinned fingerprint openSPRES).")
log("")

src_parts = zparts(SRC)
src_counts = counts_by_class(src_parts)
s0 = pptx_stats(SRC)
log(f"SOURCE: {len(src_parts)} zip parts; classes {sorted(src_counts.items())}")
log(f"SOURCE pptx stats: {s0}")
log("")

# T0: open -> save unchanged
prs = Presentation(SRC)
prs.save(T0)
t0_parts = zparts(T0)
t0_counts = counts_by_class(t0_parts)
s_t0 = pptx_stats(T0)
lost0 = sorted(set(src_parts) - set(t0_parts))
gain0 = sorted(set(t0_parts) - set(src_parts))
log(f"T0 (open->save): {len(t0_parts)} parts; classes {sorted(t0_counts.items())}")
log(f"T0 stats: {s_t0}")
log(f"T0 parts LOST vs source ({len(lost0)}): {lost0}")
log(f"T0 parts GAINED vs source ({len(gain0)}): {gain0}")
log("")

# T1: open -> add one slide on an existing layout -> save
prs = Presentation(SRC)
layout_names = [(mi, li, ly.name) for mi, m in enumerate(prs.slide_masters)
                for li, ly in enumerate(m.slide_layouts)]
log(f"T1 layouts available ({len(layout_names)}): {layout_names}")
# pick a title+content-like layout: prefer one whose name hints body,
# else layout index 1 of master 0 (convention), else 0
pick = None
for mi, li, nm in layout_names:
    if any(k in nm.lower() for k in ("title and content", "titolo e contenuto", "content")):
        pick = (mi, li, nm)
        break
if pick is None:
    m0 = prs.slide_masters[0]
    li = 1 if len(m0.slide_layouts) > 1 else 0
    pick = (0, li, m0.slide_layouts[li].name)
mi, li, nm = pick
layout = prs.slide_masters[mi].slide_layouts[li]
slide = prs.slides.add_slide(layout)
got_title = False
for ph in slide.placeholders:
    if ph.placeholder_format.idx == 0:
        ph.text = "S-PRES round-trip probe slide"
        got_title = True
log(f"T1: added slide on layout [{mi}][{li}] '{nm}'; title set: {got_title}")
prs.save(T1)
t1_parts = zparts(T1)
t1_counts = counts_by_class(t1_parts)
s_t1 = pptx_stats(T1)
lost1 = sorted(set(src_parts) - set(t1_parts))
gain1 = sorted(set(t1_parts) - set(src_parts))
log(f"T1 (open->add->save): {len(t1_parts)} parts; classes {sorted(t1_counts.items())}")
log(f"T1 stats: {s_t1}")
log(f"T1 parts LOST vs source ({len(lost1)}): {lost1}")
log(f"T1 parts GAINED vs source ({len(gain1)}): {gain1}")
log("")

# Verdict
CONTENT = {"slide", "media", "master", "layout", "theme", "notes", "chart", "embedding"}
def content_lost(lost):
    return [n for n in lost if classify(n) in CONTENT]

viol = []
cl0 = content_lost(lost0)
cl1 = content_lost(lost1)
if cl0:
    viol.append(f"T0 content parts lost: {cl0}")
if cl1:
    viol.append(f"T1 content parts lost: {cl1}")
if s_t0["n_slides"] != s0["n_slides"]:
    viol.append("T0 slide count changed")
if s_t1["n_slides"] != s0["n_slides"] + 1:
    viol.append("T1 slide count != source+1")
for k in ("n_masters", "n_layouts", "w_emu", "h_emu"):
    if s_t0[k] != s0[k]:
        viol.append(f"T0 {k} changed")
    if s_t1[k] != s0[k]:
        viol.append(f"T1 {k} changed")
mcls0 = {k: t0_counts.get(k, 0) for k in CONTENT}
mcls_src = {k: src_counts.get(k, 0) for k in CONTENT}
if mcls0 != mcls_src:
    viol.append(f"T0 content class counts differ: src={mcls_src} t0={mcls0}")

log("## VERDICT")
if viol:
    log("FAIL — violations:")
    for v in viol:
        log(f"- {v}")
    verdict = 1
else:
    noncontent_lost = [n for n in lost0 + lost1 if classify(n) not in CONTENT]
    log("PASS — no content part lost, counts preserved, extension slide landed.")
    if noncontent_lost:
        log(f"Declared non-content churn (accessory parts): {sorted(set(noncontent_lost))}")
    verdict = 0

with io.open(REPORT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines) + "\n")
print("report written; verdict:", "FAIL" if verdict else "PASS")
sys.exit(verdict)
